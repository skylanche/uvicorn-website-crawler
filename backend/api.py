"""
FastAPI server — local backend for the link crawler
Run with: uvicorn api:app --reload --port 8000
"""

from fastapi import FastAPI, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import sqlite3, uuid, os, csv, io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from crawler import start_crawl, init_db, get_meta

app = FastAPI(title="Link Crawler API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

JOBS_DIR = "jobs"
EXPORT_DIR = "exports"
os.makedirs(JOBS_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)


# ── Models ────────────────────────────────────────────────────────────────────

class CrawlRequest(BaseModel):
    url: str
    max_pages: int = 1000
    concurrency: int = 20
    delay: float = 0.0
    own_site: bool = True


# ── Helpers ───────────────────────────────────────────────────────────────────

def db_path(job_id):
    return os.path.join(JOBS_DIR, f"{job_id}.db")


def get_job_status(job_id):
    path = db_path(job_id)
    if not os.path.exists(path):
        return None
    conn = sqlite3.connect(path)
    meta = {
        "status":        get_meta(conn, "status") or "queued",
        "pages_crawled": int(get_meta(conn, "pages_crawled") or 0),
        "links_found":   int(get_meta(conn, "links_found") or 0),
        "current_url":   get_meta(conn, "current_url") or "",
        "start_url":     get_meta(conn, "start_url") or "",
        "started_at":    get_meta(conn, "started_at") or "",
        "finished_at":   get_meta(conn, "finished_at") or "",
    }
    conn.close()
    return meta


def status_color(status):
    try:
        code = int(status)
        if 200 <= code < 300: return "C6EFCE"
        if 300 <= code < 400: return "FFEB9C"
        if 400 <= code < 500: return "FFC7CE"
        if code >= 500:       return "FF0000"
    except: pass
    return "D9D9D9"


def build_xlsx(job_id):
    path = db_path(job_id)
    conn = sqlite3.connect(path)
    df = pd.read_sql("SELECT from_page, link_url, link_text, link_type, status_code, final_url FROM links", conn)
    conn.close()

    wb = openpyxl.Workbook()
    hdr_font  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    body_font = Font(name="Arial", size=10)
    hdr_fill  = PatternFill("solid", start_color="1a1a2e")
    center    = Alignment(horizontal="center")
    col_widths = [65, 70, 40, 12, 14, 65]

    def write_sheet(ws, dataframe, title_fill=None):
        fill = title_fill or hdr_fill
        for ci, col in enumerate(dataframe.columns, 1):
            c = ws.cell(1, ci, col)
            c.font = hdr_font
            c.fill = fill
            c.alignment = center
        for ri, row in enumerate(dataframe.itertuples(index=False), 2):
            for ci, val in enumerate(row, 1):
                c = ws.cell(ri, ci, str(val) if val else "")
                c.font = body_font
                if dataframe.columns[ci-1] == "status_code":
                    c.fill = PatternFill("solid", start_color=status_color(val))
                    c.alignment = center
        for i, w in enumerate(col_widths[:dataframe.shape[1]], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if dataframe.shape[0] > 0:
            ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws0 = wb.active
    ws0.title = "Summary"
    total = len(df)
    internal = len(df[df.link_type == "internal"])
    external = len(df[df.link_type == "external"])
    def cnt(lo, hi):
        return len(df[df.status_code.apply(lambda x: lo <= int(x) < hi if str(x).isdigit() else False)])
    rows = [
        ["Metric", "Value"],
        ["Total Links", total],
        ["Internal Links", internal],
        ["External Links", external],
        ["", ""],
        ["2xx OK", cnt(200,300)],
        ["3xx Redirects", cnt(300,400)],
        ["4xx Errors", cnt(400,500)],
        ["5xx Errors", cnt(500,600)],
        ["Connection Errors", len(df[~df.status_code.astype(str).str.match(r'^\d+$')])],
        ["", ""],
        ["Crawl Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            c = ws0.cell(ri, ci, val)
            c.font = hdr_font if ri == 1 else body_font
            if ri == 1: c.fill = hdr_fill; c.alignment = center
    ws0.column_dimensions["A"].width = 25
    ws0.column_dimensions["B"].width = 20

    # All links
    ws1 = wb.create_sheet("All Links")
    write_sheet(ws1, df)

    # Internal
    ws2 = wb.create_sheet("Internal Links")
    write_sheet(ws2, df[df.link_type == "internal"].reset_index(drop=True))

    # External
    ws3 = wb.create_sheet("External Links")
    write_sheet(ws3, df[df.link_type == "external"].reset_index(drop=True))

    # Broken
    ws4 = wb.create_sheet("Broken Links")
    broken = df[df.status_code.apply(
        lambda x: (str(x).isdigit() and int(x) >= 400) or not str(x).isdigit()
    )].reset_index(drop=True)
    write_sheet(ws4, broken, PatternFill("solid", start_color="C0392B"))

    out = os.path.join(EXPORT_DIR, f"{job_id}.xlsx")
    wb.save(out)
    return out


def build_csv(job_id):
    path = db_path(job_id)
    conn = sqlite3.connect(path)
    df = pd.read_sql("SELECT from_page, link_url, link_text, link_type, status_code, final_url FROM links", conn)
    conn.close()
    out = os.path.join(EXPORT_DIR, f"{job_id}.csv")
    df.to_csv(out, index=False)
    return out


# ── Routes ────────────────────────────────────────────────────────────────────

@app.post("/crawl")
def start(req: CrawlRequest, bg: BackgroundTasks):
    if not req.url.startswith(("http://", "https://")):
        req.url = "https://" + req.url
    job_id = str(uuid.uuid4())
    # Create DB immediately so status endpoint works right away
    conn = sqlite3.connect(db_path(job_id))
    conn.execute("CREATE TABLE IF NOT EXISTS job_meta (key TEXT PRIMARY KEY, value TEXT)")
    conn.execute("INSERT OR REPLACE INTO job_meta VALUES ('status','queued')")
    conn.execute(f"INSERT OR REPLACE INTO job_meta VALUES ('start_url','{req.url}')")
    conn.commit()
    conn.close()
    bg.add_task(start_crawl, job_id, req.url, req.max_pages,
                req.concurrency, req.delay)
    return {"job_id": job_id}


@app.get("/status/{job_id}")
def status(job_id: str):
    data = get_job_status(job_id)
    if not data:
        return JSONResponse({"error": "Job not found"}, 404)
    return data


@app.get("/jobs")
def list_jobs():
    jobs = []
    for f in os.listdir(JOBS_DIR):
        if f.endswith(".db"):
            job_id = f[:-3]
            data = get_job_status(job_id)
            if data:
                data["job_id"] = job_id
                jobs.append(data)
    return sorted(jobs, key=lambda x: x.get("started_at",""), reverse=True)


@app.get("/preview/{job_id}")
def preview(job_id: str, limit: int = 100, link_type: str = "all", status_filter: str = "all"):
    path = db_path(job_id)
    if not os.path.exists(path):
        return JSONResponse({"error": "Not found"}, 404)
    conn = sqlite3.connect(path)
    query = "SELECT from_page, link_url, link_text, link_type, status_code, final_url FROM links WHERE 1=1"
    params = []
    if link_type != "all":
        query += " AND link_type=?"; params.append(link_type)
    if status_filter == "broken":
        query += " AND (CAST(status_code AS INTEGER) >= 400 OR status_code NOT GLOB '[0-9]*')"
    elif status_filter == "ok":
        query += " AND CAST(status_code AS INTEGER) BETWEEN 200 AND 299"
    elif status_filter == "redirect":
        query += " AND CAST(status_code AS INTEGER) BETWEEN 300 AND 399"
    query += f" LIMIT {limit}"
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return {"rows": rows, "columns": ["From Page","Link URL","Link Text","Type","Status","Final URL"]}


@app.get("/download/{job_id}/xlsx")
def download_xlsx(job_id: str):
    out = build_xlsx(job_id)
    return FileResponse(out, filename="crawl_results.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/download/{job_id}/csv")
def download_csv(job_id: str):
    out = build_csv(job_id)
    return FileResponse(out, filename="crawl_results.csv", media_type="text/csv")


@app.delete("/jobs/{job_id}")
def delete_job(job_id: str):
    path = db_path(job_id)
    if os.path.exists(path):
        os.remove(path)
    for ext in ["xlsx", "csv"]:
        ep = os.path.join(EXPORT_DIR, f"{job_id}.{ext}")
        if os.path.exists(ep):
            os.remove(ep)
    return {"deleted": job_id}


# Serve frontend
if os.path.exists("../frontend"):
    app.mount("/", StaticFiles(directory="../frontend", html=True), name="frontend")
